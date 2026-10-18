import sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# --------------------------------------------------
# DATABASE
# --------------------------------------------------

connection = sqlite3.connect("gwcbot.db")
cursor = connection.cursor()


# --------------------------------------------------
# EXCEL WORKBOOK
# --------------------------------------------------

workbook = Workbook()

default_sheet = workbook.active
workbook.remove(default_sheet)


# --------------------------------------------------
# BRAND COLORS
# --------------------------------------------------

PRIMARY_BLUE = "5B78C7"
DARK_BLUE    = "3F5FAF"
LIGHT_BLUE   = "EAF1FF"
PALE_BLUE    = "DCE7FF"
WHITE        = "FFFFFF"
TEXT_DARK    = "2E3F73"
BORDER_BLUE  = "9CB4E8"


# --------------------------------------------------
# STYLES
# --------------------------------------------------

header_fill = PatternFill(
    fill_type="solid",
    start_color=DARK_BLUE,
    end_color=DARK_BLUE
)

title_fill = PatternFill(
    fill_type="solid",
    start_color=PALE_BLUE,
    end_color=PALE_BLUE
)

alternate_fill = PatternFill(
    fill_type="solid",
    start_color=LIGHT_BLUE,
    end_color=LIGHT_BLUE
)

thin_border = Border(
    left=Side(style="thin", color=BORDER_BLUE),
    right=Side(style="thin", color=BORDER_BLUE),
    top=Side(style="thin", color=BORDER_BLUE),
    bottom=Side(style="thin", color=BORDER_BLUE)
)

header_font = Font(
    bold=True,
    color=WHITE,
    name="Calibri",
    size=11
)

title_font = Font(
    bold=True,
    color=TEXT_DARK,
    name="Calibri",
    size=14
)

body_font = Font(
    color=TEXT_DARK,
    name="Calibri",
    size=11
)


# --------------------------------------------------
# FUNCTION TO EXPORT A QUERY TO A SHEET
# --------------------------------------------------

def export_sheet(sheet_name, query, tab_color=PRIMARY_BLUE):

    cursor.execute(query)
    rows = cursor.fetchall()

    column_names = [
        description[0]
        for description in cursor.description
    ]

    sheet = workbook.create_sheet(title=sheet_name)
    sheet.sheet_properties.tabColor = tab_color

    # -------------------------------
    # Title row
    # -------------------------------
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(column_names)
    )

    title_cell = sheet.cell(
        row=1,
        column=1,
        value=f"GWC Bot Data Export — {sheet_name}"
    )
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    title_cell.border = thin_border

    # -------------------------------
    # Header row
    # -------------------------------
    for column_number, column_name in enumerate(column_names, start=1):
        cell = sheet.cell(
            row=2,
            column=column_number,
            value=column_name
        )

        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = thin_border

    # -------------------------------
    # Data rows
    # -------------------------------
    for row_number, row in enumerate(rows, start=3):
        for column_number, value in enumerate(row, start=1):
            cell = sheet.cell(
                row=row_number,
                column=column_number,
                value=value
            )

            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True
            )
            cell.border = thin_border

            # alternate row color
            if row_number % 2 == 1:
                cell.fill = alternate_fill

    # -------------------------------
    # Freeze panes
    # -------------------------------
    sheet.freeze_panes = "A3"

    # -------------------------------
    # Filter row
    # -------------------------------
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(column_names))}{sheet.max_row}"

    # -------------------------------
    # Better row heights
    # -------------------------------
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 22

    # -------------------------------
    # Auto-size columns
    # -------------------------------
    for column in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value is not None:
                value_length = len(str(cell.value))
                if value_length > max_length:
                    max_length = value_length

        sheet.column_dimensions[column_letter].width = min(max_length + 3, 35)


# --------------------------------------------------
# SHEET 1 — MEMBERS
# --------------------------------------------------

export_sheet(
    "Members",
    """
    SELECT
        member_id AS "Member ID",
        member_role AS "Role",
        discord_id AS "Discord ID",
        full_name AS "Full Name",
        member_dob AS "Date of Birth",
        ucf_email AS "Email",  
        major AS "Major",
        phone_number AS "Phone Number",
        member_bio AS "Bio", 
        additional_info AS "Additional Info",
        member_status AS "Status",
        points_balance AS "Points Balance"
    FROM Members
    """,
    tab_color=PRIMARY_BLUE
)


# --------------------------------------------------
# SHEET 2 — EVENTS
# --------------------------------------------------

export_sheet(
    "Events",
    """
    SELECT
        event_id,
        event_name,
        event_organizer,
        event_points
    FROM Events
    """,
    tab_color="7FA6F5"
)


# --------------------------------------------------
# SHEET 3 — EVENT PARTICIPATION
# --------------------------------------------------

export_sheet(
    "Event Participation",
    """
    SELECT
        Event_Participation.participation_id,
        Members.member_id,
        Members.full_name,
        Events.event_id,
        Events.event_name,
        Event_Participation.member_attendance,
        Events.event_points
    FROM Event_Participation

    JOIN Members
        ON Event_Participation.member_id = Members.member_id

    JOIN Events
        ON Event_Participation.event_id = Events.event_id
    """,
    tab_color="9ABAF7"
)


# --------------------------------------------------
# SAVE EXCEL FILE
# --------------------------------------------------

workbook.save("gwcbot_data.xlsx")
connection.close()

print("Database exported successfully to gwcbot_data.xlsx")